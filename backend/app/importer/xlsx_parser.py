"""Parseo de archivos .xlsx con openpyxl.

Normaliza encabezados y valores, y detecta filas inválidas o duplicadas
dentro del archivo. NO contiene reglas de negocio de conciliación.
"""

import unicodedata
from io import BytesIO

from openpyxl import load_workbook

from app.domain.enums import CategoriaRegistro
from app.importer.base import ErrorImportacion, FileParser, FilaImportada

HEADERS_CLIENTE = {
    "numero de cliente",
    "nro cliente",
    "nro. cliente",
    "nro",
    "numero",
    "cliente",
    "n de cliente",
    "codigo cliente",
}
HEADERS_TIPO = {
    "tipo de operacion",
    "tipo operacion",
    "tipo",
    "operacion",
    "tipo de gestion",
}


def _normalizar(texto) -> str:
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto).strip())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return "".join(c for c in s.lower() if c.isalnum())


HEADERS_CLIENTE_NORM = {_normalizar(h) for h in HEADERS_CLIENTE}
HEADERS_TIPO_NORM = {_normalizar(h) for h in HEADERS_TIPO}


def _celda_numero(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


class XlsxParser(FileParser):
    def parse(self, contenido: bytes, tipos_validos: dict) -> list[FilaImportada]:
        try:
            wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise ErrorImportacion(
                "No se pudo leer el archivo. ¿Es un .xlsx válido?"
            )

        ws = wb.active
        if ws is None:
            raise ErrorImportacion("El archivo no contiene hojas de cálculo.")

        filas = []
        for fila_raw in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() != "" for v in fila_raw):
                filas.append(fila_raw)

        if len(filas) < 2:
            raise ErrorImportacion("El archivo no contiene registros (requiere encabezado + filas).")

        encabezado = filas[0]
        idx_cliente, idx_tipo = self._detectar_columnas(encabezado)
        vistos: set[str] = set()
        resultados: list[FilaImportada] = []

        tipos_key = {}
        for codigo, data in tipos_validos.items():
            tipos_key[_normalizar(codigo)] = codigo
            tipos_key[_normalizar(data["nombre"])] = codigo

        for n, fila_raw in enumerate(filas[1:], start=2):
            numero = _celda_numero(fila_raw[idx_cliente]) if idx_cliente is not None else None
            tipo_raw = str(fila_raw[idx_tipo]).strip() if idx_tipo is not None and fila_raw[idx_tipo] is not None else ""

            if not numero:
                resultados.append(
                    FilaImportada(
                        fila=n,
                        cliente_numero=None,
                        tipo_archivo=_normalizar(tipo_raw) or None,
                        tipo_codigo=None,
                        categoria=CategoriaRegistro.INVALIDO.value,
                        mensaje="Número de cliente faltante o vacío.",
                    )
                )
                continue

            tipo_archivo = tipo_raw.upper().strip()
            tipo_codigo = tipos_key.get(_normalizar(tipo_raw))
            if tipo_codigo is None:
                resultados.append(
                    FilaImportada(
                        fila=n,
                        cliente_numero=numero,
                        tipo_archivo=tipo_archivo or None,
                        tipo_codigo=None,
                        categoria=CategoriaRegistro.INVALIDO.value,
                        mensaje=f"Tipo de operación no reconocido: '{tipo_raw}'.",
                    )
                )
                continue

            if numero in vistos:
                resultados.append(
                    FilaImportada(
                        fila=n,
                        cliente_numero=numero,
                        tipo_archivo=tipo_archivo,
                        tipo_codigo=tipo_codigo,
                        categoria=CategoriaRegistro.DUPLICADO.value,
                        mensaje=f"Cliente '{numero}' duplicado dentro del archivo.",
                    )
                )
                continue

            vistos.add(numero)
            resultados.append(
                FilaImportada(
                    fila=n,
                    cliente_numero=numero,
                    tipo_archivo=tipo_archivo,
                    tipo_codigo=tipo_codigo,
                    categoria=None,
                    mensaje=None,
                )
            )

        return resultados

    def _detectar_columnas(self, encabezado) -> tuple[int | None, int | None]:
        idx_cliente = idx_tipo = None
        for i, celda in enumerate(encabezado):
            key = _normalizar(celda)
            if not key:
                continue
            if idx_cliente is None and key in HEADERS_CLIENTE_NORM:
                idx_cliente = i
            elif idx_tipo is None and key in HEADERS_TIPO_NORM:
                idx_tipo = i
        if idx_cliente is None:
            raise ErrorImportacion(
                "No se encontró la columna 'número de cliente' en el encabezado."
            )
        if idx_tipo is None:
            raise ErrorImportacion(
                "No se encontró la columna 'tipo de operación' en el encabezado."
            )
        return idx_cliente, idx_tipo