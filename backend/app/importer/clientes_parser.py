"""Parseo de archivos .xlsx de carga por operador (solo números de cliente).

El encabezado es opcional: si la primera fila parece un encabezado de
clientes se ignora. De cada fila se toma la primera celda no vacía.
"""

from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

from app.importer.base import ErrorImportacion
from app.importer.xlsx_parser import HEADERS_CLIENTE_NORM, _celda_numero, _normalizar


@dataclass
class ResultadoClientes:
    clientes: list[str]
    total_filas: int
    duplicados: int
    invalidas: int


def parse_clientes(contenido: bytes) -> ResultadoClientes:
    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorImportacion(
            "No se pudo leer el archivo. ¿Es un .xlsx válido?"
        )

    ws = wb.active
    if ws is None:
        raise ErrorImportacion("El archivo no contiene hojas de cálculo.")

    filas = []
    for fila_raw in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() != "" for v in fila_raw):
            filas.append(fila_raw)

    if len(filas) < 1:
        raise ErrorImportacion("El archivo está vacío.")

    primero = _celda_numero(filas[0][0] if filas[0] else None)
    inicio = 1 if (primero is None or _normalizar(primero) in HEADERS_CLIENTE_NORM) else 0

    vistos: list[str] = []
    vistos_set: set[str] = set()
    duplicados = invalidas = 0
    for fila_raw in filas[inicio:]:
        numero = None
        for valor in fila_raw:
            if valor is not None and str(valor).strip() != "":
                numero = _celda_numero(valor)
                break
        if numero is None:
            invalidas += 1
            continue
        if numero in vistos_set:
            duplicados += 1
            continue
        vistos_set.add(numero)
        vistos.append(numero)

    if not vistos:
        raise ErrorImportacion("El archivo no contiene números de cliente válidos.")

    return ResultadoClientes(
        clientes=vistos,
        total_filas=len(filas) - inicio,
        duplicados=duplicados,
        invalidas=invalidas,
    )